"""
SLIDES

"Ultron, monta uma apresentação sobre X" → um .pptx que abre no
PowerPoint, no Google Slides e no LibreOffice.

Duas decisões que separam isto de um gerador de slide feio:

**O conteúdo vem antes do desenho.** O modelo escreve a estrutura — tema
de cada slide, os pontos, a fala do apresentador — com esquema validado.
Só depois isso vira arquivo.

**Nada de layout padrão do PowerPoint.** Os placeholders prontos geram
aquele slide com título gigante e marcador redondo que todo mundo
reconhece. Aqui cada slide é desenhado: caixa de texto posicionada,
hierarquia de tamanho, uma barra de acento e número de página.

Regra de conteúdo: no máximo 6 pontos por slide, cada um com no máximo
uma linha. Slide não é documento — o que precisa de parágrafo vai para a
nota do apresentador, que o `python-pptx` guarda e o PowerPoint mostra.
"""
from __future__ import annotations

import re
import unicodedata
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt
from pydantic import BaseModel, Field

from nucleo.config import SAIDA
from nucleo.modelos import pede_json
from nucleo.permissao import LIVRE
from . import Contexto, ferramenta

LARGURA = Inches(13.333)      # 16:9
ALTURA = Inches(7.5)
MARGEM = Inches(0.9)

PALETAS = {
    'escuro':  dict(fundo='101114', texto='F5F5F7', fraco='A1A1AA', acento='34D399'),
    'claro':   dict(fundo='FFFFFF', texto='18181B', fraco='52525B', acento='047857'),
    'azul':    dict(fundo='0B1220', texto='F8FAFC', fraco='94A3B8', acento='38BDF8'),
    'quente':  dict(fundo='1A1210', texto='FDF8F3', fraco='BFA99A', acento='F59E0B'),
}


class Slide(BaseModel):
    titulo: str = Field(description='O tema do slide, curto.')
    pontos: list[str] = Field(description='Até 6 pontos, cada um com no máximo uma linha.')
    nota: str = Field(default='', description='O que o apresentador fala. Aqui pode ter texto.')


class Apresentacao(BaseModel):
    titulo: str
    subtitulo: str = ''
    slides: list[Slide]
    fechamento: str = Field(default='', description='A última frase, a que fica.')


INSTRUCAO = """Você escreve a estrutura de uma apresentação.

Regras que valem mais que o conteúdo:
- no máximo 6 pontos por slide, cada ponto com no MÁXIMO uma linha \
(umas 12 palavras). Slide não é documento;
- o que precisa de explicação vai para `nota`, não para o slide;
- o título do slide é uma AFIRMAÇÃO, não um rótulo: "O custo dobrou em dois anos" \
em vez de "Custos";
- nada de "Introdução", "Agenda", "Obrigado" — slide que não diz nada ocupa lugar;
- se te deram material, use só o que está nele. Não invente número, data, \
nome nem citação. O que faltar, escreva como pergunta na `nota`.

Português do Brasil, direto, sem jargão de consultoria."""


def _apelido(t: str) -> str:
    p = unicodedata.normalize('NFKD', t).encode('ascii', 'ignore').decode()
    return re.sub(r'[^a-z0-9]+', '-', p.lower()).strip('-')[:50] or 'apresentacao'


def _cor(h: str) -> RGBColor:
    return RGBColor.from_string(h)


def _fundo(slide, cor: str) -> None:
    f = slide.background.fill
    f.solid()
    f.fore_color.rgb = _cor(cor)


def _caixa(slide, x, y, w, h, texto, tamanho, cor, negrito=False,
           alinhamento=PP_ALIGN.LEFT, entrelinha=1.15):
    cx = slide.shapes.add_textbox(x, y, w, h)
    tf = cx.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = alinhamento
    p.line_spacing = entrelinha
    r = p.add_run()
    r.text = texto
    r.font.size = Pt(tamanho)
    r.font.bold = negrito
    r.font.color.rgb = _cor(cor)
    r.font.name = 'Segoe UI'
    return cx


def _barra(slide, cor: str, x=MARGEM, y=Inches(1.45), w=Inches(1.1), h=Emu(45720)):
    forma = slide.shapes.add_shape(1, x, y, w, h)     # 1 = retângulo
    forma.fill.solid()
    forma.fill.fore_color.rgb = _cor(cor)
    forma.line.fill.background()
    forma.shadow.inherit = False
    return forma


def monta(a: Apresentacao, destino: Path, tema: str = 'escuro') -> Path:
    cores = PALETAS.get(tema, PALETAS['escuro'])
    pres = Presentation()
    pres.slide_width, pres.slide_height = LARGURA, ALTURA
    vazio = pres.slide_layouts[6]

    # Capa
    s = pres.slides.add_slide(vazio)
    _fundo(s, cores['fundo'])
    _barra(s, cores['acento'], y=Inches(2.5), w=Inches(1.6))
    _caixa(s, MARGEM, Inches(2.9), LARGURA - 2 * MARGEM, Inches(2.0),
           a.titulo, 46, cores['texto'], negrito=True, entrelinha=1.05)
    if a.subtitulo:
        _caixa(s, MARGEM, Inches(4.5), LARGURA - 2 * MARGEM, Inches(1.0),
               a.subtitulo, 20, cores['fraco'])

    # Conteúdo
    for i, sl in enumerate(a.slides, 1):
        s = pres.slides.add_slide(vazio)
        _fundo(s, cores['fundo'])
        _caixa(s, MARGEM, Inches(0.7), LARGURA - 2 * MARGEM, Inches(1.1),
               sl.titulo, 30, cores['texto'], negrito=True, entrelinha=1.1)
        _barra(s, cores['acento'])
        topo = Inches(2.0)
        for ponto in sl.pontos[:6]:
            _caixa(s, MARGEM + Inches(0.35), topo, LARGURA - 2 * MARGEM - Inches(0.35),
                   Inches(0.6), '—  ' + ponto.strip(), 18, cores['texto'])
            topo += Inches(0.72)
        _caixa(s, LARGURA - MARGEM - Inches(1.2), ALTURA - Inches(0.75), Inches(1.2),
               Inches(0.4), str(i), 12, cores['fraco'], alinhamento=PP_ALIGN.RIGHT)
        if sl.nota:
            s.notes_slide.notes_text_frame.text = sl.nota

    # Fechamento
    if a.fechamento:
        s = pres.slides.add_slide(vazio)
        _fundo(s, cores['fundo'])
        _barra(s, cores['acento'], y=Inches(3.0), w=Inches(1.6))
        _caixa(s, MARGEM, Inches(3.4), LARGURA - 2 * MARGEM, Inches(2.0),
               a.fechamento, 34, cores['texto'], negrito=True, entrelinha=1.15)

    destino.parent.mkdir(parents=True, exist_ok=True)
    pres.save(str(destino))
    return destino


class SlidesArgs(BaseModel):
    assunto: str = Field(description='Sobre o que é a apresentação, e para quem.')
    quantos: int = Field(default=8, description='Quantos slides de conteúdo.')
    material: str = Field(default='', description='Texto de base, se houver. Ou um caminho '
                                                 'de arquivo para ele ler.')
    tema: str = Field(default='escuro', description='escuro | claro | azul | quente')
    salvar_em: str = Field(default='', description='Caminho do .pptx. Vazio = pasta saida/.')
    abrir: bool = Field(default=True, description='Abrir depois de montar.')


@ferramenta('criar_slides',
            'Monta uma apresentação em PowerPoint (.pptx) sobre um assunto, '
            'com nota do apresentador em cada slide.',
            SlidesArgs, nivel=LIVRE, resumo=lambda a: f'montar {a.quantos} slides sobre {a.assunto[:70]}')
def criar_slides(a: SlidesArgs, ctx: Contexto) -> str:
    material = a.material
    if material and len(material) < 400 and Path(material).expanduser().exists():
        material = Path(material).expanduser().read_text(encoding='utf-8', errors='replace')[:60000]

    pedido = f'Assunto: {a.assunto}\nQuantidade de slides de conteúdo: {a.quantos}'
    if material:
        pedido += f'\n\nMaterial de base (use só o que está aqui):\n{material}'

    try:
        estrutura: Apresentacao = pede_json(INSTRUCAO, pedido, Apresentacao, ctx.cfg.modelo,
                                            max_tokens=12000)
    except Exception as e:
        return f'não consegui montar a estrutura: {e}'

    destino = (Path(a.salvar_em).expanduser() if a.salvar_em
               else SAIDA / 'slides' / f'{_apelido(estrutura.titulo)}.pptx')
    try:
        monta(estrutura, destino, a.tema)
    except Exception as e:
        return f'a estrutura ficou pronta mas o arquivo falhou: {e}'

    if a.abrir:
        try:
            from .computador import AbrirArgs, abrir_programa
            abrir_programa(AbrirArgs(programa=str(destino)), ctx)
        except Exception:
            pass
    return (f'pronto: {destino}\n'
            f'"{estrutura.titulo}" — {len(estrutura.slides)} slides, '
            'com nota do apresentador em cada um.')
