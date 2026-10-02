#!/usr/bin/env python3
"""
TESTES — python3 testes.py

Nada aqui toca a rede, o microfone ou a sua máquina de verdade: a API do
Claude entra como dublê, os arquivos vão para uma pasta temporária e o
navegador não abre.

O que está sendo testado não é "o código roda" — é o que decide se este
projeto é seguro de usar:

  • voz sozinha NUNCA autoriza coisa irreversível;
  • 'sempre' vale para o que é reversível, nunca para o que não é;
  • permissão negada vira resposta honesta para o modelo, não outro caminho;
  • dinheiro não é contado duas vezes, e não se inventa centavo;
  • cortar conversa antiga não quebra o par ferramenta/resultado.
"""
from __future__ import annotations

import contextlib
import io
import shutil
import sys
import tempfile
import time
import traceback
import types
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from ferramentas import REGISTRO, Contexto, carrega_tudo, catalogo      # noqa: E402
from nucleo import config                                              # noqa: E402
from nucleo.cerebro import Cerebro                                     # noqa: E402
from nucleo.permissao import (CUIDADO, LIVRE, PERIGO, Porteiro,        # noqa: E402
                              Veredito, avalia_caminho, avalia_comando)
from nucleo.registro import Diario                                     # noqa: E402
from nucleo.vendas import Caixa, centavos, inicio_do_dia, reais        # noqa: E402
from nucleo.voz import limpa_para_fala                                 # noqa: E402

CASOS: list[tuple[str, object]] = []
TMP = Path(tempfile.mkdtemp(prefix='jarvis-testes-'))
carrega_tudo()


def teste(nome):
    def envolve(f):
        CASOS.append((nome, f))
        return f
    return envolve


def igual(a, b, o_que=''):
    if a != b:
        raise AssertionError(f'{o_que or "valor"}: esperava {b!r}, veio {a!r}')


def verdade(x, o_que=''):
    if not x:
        raise AssertionError(o_que or 'esperava verdadeiro')


def cfg_teste(**extra):
    c = config.Config(anthropic='x', raizes_seguras=[str(TMP / 'livre')])
    (TMP / 'livre').mkdir(parents=True, exist_ok=True)
    for k, v in extra.items():
        setattr(c, k, v)
    return c


# ══ dublê do Claude ═════════════════════════════════════════════════
def bloco_texto(t):
    return types.SimpleNamespace(type='text', text=t)


def bloco_uso(nome, entrada, ident='u1'):
    return types.SimpleNamespace(type='tool_use', name=nome, input=entrada, id=ident)


class FalsoFluxo:
    def __init__(self, msg):
        self.msg = msg

    def __enter__(self): return self
    def __exit__(self, *_): return False

    def __iter__(self):
        for b in self.msg.content:
            if getattr(b, 'type', '') == 'text':
                # Em pedaços, como a API manda de verdade.
                for pedaco in _fatia(b.text):
                    yield types.SimpleNamespace(type='text', text=pedaco)

    def get_final_message(self):
        return self.msg


def _fatia(t, n=9):
    return [t[i:i + n] for i in range(0, len(t), n)] or ['']


class FalsoClaude:
    """Devolve as respostas na ordem em que foram programadas."""

    def __init__(self, *respostas):
        self.roteiro = list(respostas)
        self.chamadas: list[dict] = []
        self.messages = self

    def stream(self, **kw):
        # Cópia da lista: o cérebro continua escrevendo no MESMO histórico
        # depois desta chamada, e guardar a referência faria o teste olhar
        # o estado final em vez do estado daquele momento.
        self.chamadas.append({**kw, 'messages': list(kw['messages'])})
        blocos, parada = self.roteiro.pop(0) if self.roteiro else ([bloco_texto('ok')], 'end_turn')
        return FalsoFluxo(types.SimpleNamespace(content=blocos, stop_reason=parada))


class FalsoGemini:
    """Imita o streaming do google-genai: pedaços com partes de texto ou chamada."""

    def __init__(self, *respostas):
        self.roteiro = list(respostas)
        self.chamadas: list[dict] = []
        self.models = self

    def generate_content_stream(self, **kw):
        self.chamadas.append({**kw, 'contents': list(kw['contents'])})
        partes = self.roteiro.pop(0) if self.roteiro else [parte_texto('ok')]
        # Cada parte vira um pedaço separado, como a API manda.
        return iter([types.SimpleNamespace(
            candidates=[types.SimpleNamespace(
                content=types.SimpleNamespace(parts=[p]))]) for p in partes])


def parte_texto(t, pensamento=False):
    return types.SimpleNamespace(text=t, function_call=None, thought=pensamento)


def parte_chamada(nome, args):
    return types.SimpleNamespace(
        text=None, thought=False,
        function_call=types.SimpleNamespace(name=nome, args=args, id='c1'))


def monta_gemini(cliente, cfg=None, respostas_porteiro=None):
    from nucleo.cerebro import CerebroGemini
    cfg = cfg or cfg_teste(provedor='gemini', gemini='x')
    cfg.provedor = 'gemini'
    diario = Diario(TMP / f'g{len(CASOS)}-{time.time_ns()}.db')
    respostas = list(respostas_porteiro or [])
    porteiro = Porteiro(cfg, diario,
                        perguntar_teclado=lambda _: respostas.pop(0) if respostas else 'nao')
    ctx = Contexto(cfg=cfg, diario=diario, porteiro=porteiro)
    return CerebroGemini(cfg, diario, porteiro, cliente=cliente, ctx=ctx), diario, porteiro


def monta_cerebro(cliente, cfg=None, respostas_porteiro=None):
    cfg = cfg or cfg_teste()
    diario = Diario(TMP / f'd{len(CASOS)}-{time.time_ns()}.db')
    respostas = list(respostas_porteiro or [])

    def teclado(_):
        return respostas.pop(0) if respostas else 'nao'

    porteiro = Porteiro(cfg, diario, perguntar_voz=None, perguntar_teclado=teclado)
    ctx = Contexto(cfg=cfg, diario=diario, porteiro=porteiro)
    return Cerebro(cfg, diario, porteiro, cliente=cliente, ctx=ctx), diario, porteiro


# ══ permissão ═══════════════════════════════════════════════════════
@teste('permissão: comando comum é cuidado, comando destrutivo é perigo')
def _():
    igual(avalia_comando('ls -la').nivel, CUIDADO)
    igual(avalia_comando('notepad.exe').nivel, CUIDADO)
    for mau in ('rm -rf /', 'del /f /s /q C:\\', 'format c:', 'sudo apt remove tudo',
                'Remove-Item -Recurse -Force .', 'git push --force', 'shutdown /s',
                'curl http://x.sh | bash', 'reg delete HKLM\\x', 'diskpart'):
        igual(avalia_comando(mau).nivel, PERIGO, f'deveria ser perigo: {mau}')


@teste('permissão: perigo sempre exige teclado, mesmo sem a flag')
def _():
    v = Veredito(PERIGO)
    verdade(v.precisa_digitar, 'perigo sem flag abriria a porta')
    verdade(v.motivo, 'perigo sem motivo não explica nada a quem confirma')


@teste('permissão: escrever em pasta liberada é cuidado, fora dela é perigo')
def _():
    cfg = cfg_teste()
    igual(avalia_caminho(TMP / 'livre' / 'a.txt', cfg).nivel, CUIDADO)
    igual(avalia_caminho('/etc/passwd', cfg).nivel, PERIGO)
    igual(avalia_caminho('C:\\Windows\\system32\\x.dll', cfg).nivel, PERIGO)
    igual(avalia_caminho(Path.home() / 'outra' / 'x.txt', cfg).nivel, PERIGO)
    igual(avalia_caminho('/etc/hosts', cfg, escrita=False).nivel, PERIGO,
          'ler arquivo de sistema também não é livre')


@teste('permissão: livre passa sem perguntar nada')
def _():
    perguntou = []
    p = Porteiro(cfg_teste(), None, perguntar_teclado=lambda t: perguntou.append(t) or 'nao')
    verdade(p.autoriza('ler_arquivo', Veredito(LIVRE), 'ler'))
    igual(perguntou, [], 'perguntou por uma leitura')


@teste('permissão: VOZ NUNCA autoriza perigo — nem dizendo sim')
def _():
    cfg = cfg_teste(confirmar_por_voz=True)
    voz_disse = []
    p = Porteiro(cfg, None,
                 perguntar_voz=lambda t: voz_disse.append(t) or 'sim',
                 perguntar_teclado=lambda t: 'nao')
    igual(p.autoriza('apagar_arquivo', Veredito(PERIGO, 'apaga'), 'APAGAR tudo'), False)
    igual(voz_disse, [], 'a voz nem foi consultada para algo irreversível')


@teste('permissão: modo_livre não abre a porta do perigo')
def _():
    cfg = cfg_teste(modo_livre=True)
    p = Porteiro(cfg, None, perguntar_teclado=lambda t: 'nao')
    verdade(p.autoriza('rodar_comando', Veredito(CUIDADO), 'ls'), 'modo livre deveria passar')
    igual(p.autoriza('rodar_comando', Veredito(PERIGO, 'rm -rf'), 'rm -rf'), False)


@teste("permissão: 'sempre' vale para cuidado e não vale para perigo")
def _():
    cfg = cfg_teste(confirmar_por_voz=False)
    respostas = ['sempre']
    p = Porteiro(cfg, None, perguntar_teclado=lambda t: respostas.pop(0) if respostas else 'nao')
    verdade(p.autoriza('rodar_comando', Veredito(CUIDADO), 'ls'))
    verdade(p.autoriza('rodar_comando', Veredito(CUIDADO), 'pwd'), 'não lembrou do sempre')
    # mesmo com 'sempre' guardado, perigo pergunta de novo — e aqui diz não
    igual(p.autoriza('rodar_comando', Veredito(PERIGO, 'rm'), 'rm -rf'), False)


# ══ diário ══════════════════════════════════════════════════════════
@teste('diário: registra ação, lembra fato e conta o dia')
def _():
    d = Diario(TMP / 'diario.db')
    d.acao('rodar_comando', {'comando': 'ls'}, CUIDADO, 'feito', 'a.txt b.txt', 'lista aí')
    d.acao('apagar_arquivo', {'caminho': 'x'}, PERIGO, 'recusado', '', 'apaga x')
    r = d.resumo_do_dia()
    igual(r['total'], 2)
    d.lembra('contador', 'o contador é o Edson')
    igual(d.lembretes()['contador'], 'o contador é o Edson')
    verdade(d.esquece('contador'))
    igual(d.lembretes(), {})
    d.fechar()


# ══ dinheiro ════════════════════════════════════════════════════════
@teste('dinheiro: lê o formato brasileiro e o americano sem trocar cem vezes')
def _():
    igual(centavos('R$ 1.200,50'), 120050)
    igual(centavos('1,200.50'), 120050)
    igual(centavos('R$ 80'), 8000)
    igual(centavos('80,00'), 8000)
    igual(centavos('1.200'), 120000, 'mil e duzentos, não um e vinte')
    igual(centavos('0,99'), 99)
    igual(centavos('abc'), 0)
    igual(centavos(''), 0)
    igual(reais(120050), 'R$ 1.200,50')


@teste('dinheiro: a mesma mensagem lida duas vezes não soma duas vezes')
def _():
    c = Caixa(TMP / 'caixa1.db')
    agora = int(time.time())
    verdade(c.guarda(50000, agora, 'Padaria', 'pix', 'Grupo Vendas', 'pix de 500 feito'))
    igual(c.guarda(50000, agora, 'Padaria', 'pix', 'Grupo Vendas', 'pix de 500 feito'), False,
          'contou de novo')
    total, n = c.total(inicio_do_dia())
    igual((total, n), (50000, 1))
    c.fechar()


@teste('dinheiro: total, por cliente e marcação do que ficou duvidoso')
def _():
    c = Caixa(TMP / 'caixa2.db')
    hoje = inicio_do_dia()
    c.guarda(120000, hoje + 100, 'Grão Dourado', 'site', 'z', 'pagou 1200', 1.0)
    c.guarda(30000, hoje + 200, 'Grão Dourado', 'ajuste', 'z', 'mais 300', 1.0)
    c.guarda(9900, hoje + 300, 'Mega Express', '?', 'z', 'talvez 99', 0.4)
    igual(c.total(hoje)[0], 159900)
    igual(c.por_cliente(hoje)[0], ('Grão Dourado', 150000, 2))
    igual(len(c.duvidosas(hoje)), 1, 'o de confiança baixa precisa ficar marcado')
    c.fechar()


# ══ catálogo ════════════════════════════════════════════════════════
@teste('catálogo: todo esquema está completo e válido para a API')
def _():
    c = catalogo()
    verdade(len(c) >= 45, f'poucas ferramentas: {len(c)}')
    nomes = set()
    for e in c:
        verdade(e['name'] and e['name'] not in nomes, f'nome repetido: {e["name"]}')
        nomes.add(e['name'])
        verdade(len(e['description']) > 15, f'{e["name"]} sem descrição útil')
        verdade(e['input_schema'].get('type') == 'object', f'{e["name"]} sem esquema')
        verdade('title' not in e['input_schema'], f'{e["name"]} com title sobrando')


@teste('catálogo: toda ferramenta sabe se descrever antes de agir')
def _():
    for nome, f in REGISTRO.items():
        campos = f.args.model_fields
        obrigatorios = {k: 'x' for k, v in campos.items() if v.is_required()}
        try:
            args = f.args(**obrigatorios)
        except Exception:
            continue      # campo obrigatório que não é texto; o laço cobre
        d = f.descreve(args)
        verdade(isinstance(d, str) and d, f'{nome} não se descreve')


# ══ arquivos ════════════════════════════════════════════════════════
@teste('arquivos: escreve, lê, lista e procura')
def _():
    from ferramentas import arquivos as A
    ctx = Contexto(cfg=cfg_teste())
    alvo = TMP / 'livre' / 'nota.txt'
    A.escrever_arquivo(A.EscreverArgs(caminho=str(alvo), conteudo='linha um\nlinha dois'), ctx)
    igual(A.ler_arquivo(A.LerArgs(caminho=str(alvo)), ctx), 'linha um\nlinha dois')
    igual(A.ler_arquivo(A.LerArgs(caminho=str(alvo), linhas=1), ctx), 'linha um')
    verdade('nota.txt' in A.listar_pasta(A.ListarArgs(pasta=str(TMP / 'livre')), ctx))
    verdade('linha dois' in A.procurar_em_arquivos(
        A.ProcurarArgs(termo='linha dois', pasta=str(TMP / 'livre')), ctx))
    igual(A.ler_arquivo(A.LerArgs(caminho=str(TMP / 'nao-existe')), ctx)[:9], 'não exist')


@teste('arquivos: apagar é perigo declarado, não "cuidado"')
def _():
    igual(REGISTRO['apagar_arquivo'].nivel, PERIGO)
    v = REGISTRO['apagar_arquivo'].veredito(
        REGISTRO['apagar_arquivo'].args(caminho='x'), None)
    verdade(v.precisa_digitar)


# ══ o laço do cérebro ═══════════════════════════════════════════════
@teste('cérebro: usa a ferramenta, lê o resultado e responde')
def _():
    alvo = TMP / 'livre' / 'leia.txt'
    alvo.write_text('o segredo é 42')
    cli = FalsoClaude(
        ([bloco_uso('ler_arquivo', {'caminho': str(alvo)})], 'tool_use'),
        ([bloco_texto('O arquivo diz que o segredo é 42.')], 'end_turn'))
    c, diario, _ = monta_cerebro(cli)
    r = c.responde('o que tem no arquivo leia.txt?')
    igual(r, 'O arquivo diz que o segredo é 42.')
    # o resultado da ferramenta voltou para o modelo
    segunda = cli.chamadas[1]['messages']
    resultado = segunda[-1]['content'][0]
    igual(resultado['type'], 'tool_result')
    verdade('42' in resultado['content'])
    igual(resultado['is_error'], False)


@teste('cérebro: permissão negada vira resposta honesta, não outro caminho')
def _():
    cli = FalsoClaude(
        ([bloco_uso('rodar_comando', {'comando': 'rm -rf /'})], 'tool_use'),
        ([bloco_texto('Não fiz: você não autorizou.')], 'end_turn'))
    c, diario, _ = monta_cerebro(cli, respostas_porteiro=['nao'])
    r = c.responde('apaga tudo')
    resultado = cli.chamadas[1]['messages'][-1]['content'][0]
    verdade(resultado['is_error'], 'precisa voltar como erro')
    verdade('NÃO AUTORIZOU' in resultado['content'])
    verdade('Não tente outro caminho' in resultado['content'],
            'sem isso o modelo tenta a mesma coisa por outra ferramenta')
    igual(r, 'Não fiz: você não autorizou.')
    acoes = diario.ultimas(5)
    igual(acoes[0]['decisao'], 'recusado', 'a recusa fica registrada')


@teste('cérebro: comando perigoso só passa com confirmação digitada')
def _():
    marca = TMP / 'livre' / 'passou.txt'
    cli = FalsoClaude(
        ([bloco_uso('rodar_comando', {'comando': f'rm -f {marca} && echo feito'})], 'tool_use'),
        ([bloco_texto('feito')], 'end_turn'))
    cfg = cfg_teste(confirmar_por_voz=True)
    diario = Diario(TMP / 'perigo.db')
    voz_consultada = []
    porteiro = Porteiro(cfg, diario,
                        perguntar_voz=lambda t: voz_consultada.append(t) or 'sim',
                        perguntar_teclado=lambda t: 'sim')
    ctx = Contexto(cfg=cfg, diario=diario, porteiro=porteiro)
    c = Cerebro(cfg, diario, porteiro, cliente=cli, ctx=ctx)
    c.responde('limpa aquele arquivo')
    igual(voz_consultada, [], 'a voz foi consultada para um rm -f')
    igual(diario.ultimas(3)[0]['nivel'], PERIGO)


@teste('cérebro: ferramenta que não existe não derruba o laço')
def _():
    cli = FalsoClaude(
        ([bloco_uso('voar', {})], 'tool_use'),
        ([bloco_texto('essa eu não tenho')], 'end_turn'))
    c, _, _ = monta_cerebro(cli)
    igual(c.responde('voa'), 'essa eu não tenho')
    verdade('não existe' in cli.chamadas[1]['messages'][-1]['content'][0]['content'])


@teste('cérebro: argumento inválido volta como erro, não como exceção')
def _():
    cli = FalsoClaude(
        ([bloco_uso('ler_arquivo', {'linhas': 'muitas'})], 'tool_use'),
        ([bloco_texto('faltou o caminho')], 'end_turn'))
    c, _, _ = monta_cerebro(cli)
    igual(c.responde('lê aí'), 'faltou o caminho')
    verdade('inválidos' in cli.chamadas[1]['messages'][-1]['content'][0]['content'])


@teste('cérebro: ferramenta que explode é reportada, não escondida')
def _():
    def explode(a, ctx):
        raise RuntimeError('o disco sumiu')
    REGISTRO['ler_arquivo'].funcao, original = explode, REGISTRO['ler_arquivo'].funcao
    try:
        cli = FalsoClaude(
            ([bloco_uso('ler_arquivo', {'caminho': 'x'})], 'tool_use'),
            ([bloco_texto('deu erro no disco')], 'end_turn'))
        c, _, _ = monta_cerebro(cli)
        c.responde('lê o x')
        conteudo = cli.chamadas[1]['messages'][-1]['content'][0]['content']
        verdade('o disco sumiu' in conteudo, 'o erro real precisa chegar ao modelo')
    finally:
        REGISTRO['ler_arquivo'].funcao = original


@teste('cérebro: fala frase por frase enquanto o texto chega')
def _():
    cli = FalsoClaude(([bloco_texto('Abri o site. Já está na tela. Quer que eu role?')],
                       'end_turn'))
    c, _, _ = monta_cerebro(cli)
    ditas = []
    c.responde('abre o site', ao_falar=ditas.append)
    igual(len(ditas), 3, f'deveria falar 3 frases, falou {ditas}')
    igual(ditas[0], 'Abri o site.')


@teste('cérebro: para de insistir depois do teto de voltas')
def _():
    cli = FalsoClaude(*[([bloco_uso('listar_pasta', {'pasta': str(TMP)}, f'u{i}')], 'tool_use')
                        for i in range(40)])
    cfg = cfg_teste(voltas_maximas=5)
    c, _, _ = monta_cerebro(cli, cfg=cfg)
    r = c.responde('fica listando')
    verdade('muitas voltas' in r)
    igual(len(cli.chamadas), 5)


@teste('cérebro: cortar conversa antiga não deixa resultado órfão')
def _():
    cli = FalsoClaude(([bloco_texto('ok')], 'end_turn'))
    c, _, _ = monta_cerebro(cli)
    for i in range(30):
        c.historico.append({'role': 'user', 'content': f'p{i}'})
        c.historico.append({'role': 'assistant',
                            'content': [types.SimpleNamespace(type='tool_use', name='x',
                                                              input={}, id=f'u{i}')]})
        c.historico.append({'role': 'user',
                            'content': [{'type': 'tool_result', 'tool_use_id': f'u{i}',
                                         'content': 'ok'}]})
    c._encolhe(teto=10)
    primeiro = c.historico[0]
    conteudo = primeiro['content']
    orfao = isinstance(conteudo, list) and any(
        isinstance(b, dict) and b.get('type') == 'tool_result' for b in conteudo)
    verdade(not orfao, 'a conversa começa com um tool_result sem o tool_use — a API recusa')
    igual(primeiro['role'], 'user')


@teste('cérebro: o sistema conta a data, as pastas liberadas e o que ele lembra')
def _():
    cli = FalsoClaude(([bloco_texto('ok')], 'end_turn'))
    c, diario, _ = monta_cerebro(cli)
    diario.lembra('contador', 'o contador é o Edson')
    c.responde('oi')
    sistema = cli.chamadas[0]['system']
    verdade('Agora:' in sistema)
    verdade('Pastas liberadas' in sistema)
    verdade('Edson' in sistema, 'o que ele lembra precisa chegar ao modelo')
    verdade('Nunca diga que fez algo que você não fez' in sistema)


# ══ o laço do Gemini ════════════════════════════════════════════════
@teste('gemini: usa a ferramenta, lê o resultado e responde')
def _():
    alvo = TMP / 'livre' / 'gemini.txt'
    alvo.write_text('o segredo é 7')
    cli = FalsoGemini(
        [parte_chamada('ler_arquivo', {'caminho': str(alvo)})],
        [parte_texto('O arquivo diz que o segredo é 7.')])
    c, diario, _ = monta_gemini(cli)
    igual(c.responde('o que tem no gemini.txt?'), 'O arquivo diz que o segredo é 7.')
    # o resultado voltou como function_response
    segunda = cli.chamadas[1]['contents']
    resposta = segunda[-1].parts[0].function_response
    igual(resposta.name, 'ler_arquivo')
    verdade('7' in str(resposta.response))
    verdade('resultado' in resposta.response, 'sucesso tem que vir como resultado')


@teste('gemini: a chamada automática do SDK fica DESLIGADA')
def _():
    # Se o SDK executar a função sozinho, a permissão nunca é consultada.
    # É a linha que separa ter trava de não ter.
    cli = FalsoGemini([parte_texto('oi')])
    c, _, _ = monta_gemini(cli)
    c.responde('oi')
    cfg_enviada = cli.chamadas[0]['config']
    igual(cfg_enviada.automatic_function_calling.disable, True)
    verdade(cfg_enviada.tools, 'as ferramentas precisam ir junto')
    decls = cfg_enviada.tools[0].function_declarations
    verdade(len(decls) >= 45, f'poucas ferramentas declaradas: {len(decls)}')


@teste('gemini: permissão negada vira resposta honesta, não outro caminho')
def _():
    cli = FalsoGemini(
        [parte_chamada('rodar_comando', {'comando': 'rm -rf /'})],
        [parte_texto('Não fiz: você não autorizou.')])
    c, diario, _ = monta_gemini(cli, respostas_porteiro=['nao'])
    c.responde('apaga tudo')
    resposta = cli.chamadas[1]['contents'][-1].parts[0].function_response
    verdade('erro' in resposta.response, 'recusa tem que voltar marcada como erro')
    texto = str(resposta.response)
    verdade('NÃO AUTORIZOU' in texto)
    verdade('Não tente outro caminho' in texto)
    igual(diario.ultimas(5)[0]['decisao'], 'recusado')


@teste('gemini: a mesma trava — voz não autoriza o irreversível')
def _():
    cfg = cfg_teste(provedor='gemini', gemini='x', confirmar_por_voz=True)
    diario = Diario(TMP / f'gv{time.time_ns()}.db')
    voz_consultada = []
    porteiro = Porteiro(cfg, diario,
                        perguntar_voz=lambda t: voz_consultada.append(t) or 'sim',
                        perguntar_teclado=lambda t: 'nao')
    from nucleo.cerebro import CerebroGemini
    cli = FalsoGemini(
        [parte_chamada('apagar_arquivo', {'caminho': str(TMP / 'livre' / 'x')})],
        [parte_texto('não autorizado')])
    c = CerebroGemini(cfg, diario, porteiro, cliente=cli,
                      ctx=Contexto(cfg=cfg, diario=diario, porteiro=porteiro))
    c.responde('apaga o x')
    igual(voz_consultada, [], 'a voz foi consultada para apagar arquivo')


@teste('gemini: pensamento do modelo não é falado')
def _():
    cli = FalsoGemini([parte_texto('Deixa eu pensar aqui.', pensamento=True),
                       parte_texto('Pronto, abri.')])
    c, _, _ = monta_gemini(cli)
    ditas = []
    r = c.responde('abre', ao_falar=ditas.append)
    igual(ditas, ['Pronto, abri.'], f'falou o que não devia: {ditas}')
    igual(r, 'Pronto, abri.')


@teste('gemini: para de insistir depois do teto de voltas')
def _():
    cli = FalsoGemini(*[[parte_chamada('listar_pasta', {'pasta': str(TMP)})]
                        for _ in range(30)])
    cfg = cfg_teste(provedor='gemini', gemini='x', voltas_maximas=4)
    c, _, _ = monta_gemini(cli, cfg=cfg)
    verdade('muitas voltas' in c.responde('fica listando'))
    igual(len(cli.chamadas), 4)


@teste('gemini: todas as ferramentas viram declaração válida para a Google')
def _():
    from google.genai import types as gt
    ruins = []
    for e in catalogo():
        try:
            gt.FunctionDeclaration(name=e['name'], description=e['description'],
                                   parameters_json_schema=e['input_schema'])
        except Exception as ex:
            ruins.append(f'{e["name"]}: {ex}')
    igual(ruins, [])


@teste('config: o provedor sai da chave que existe, sem precisar configurar')
def _():
    from nucleo.config import Config
    igual(config.Config(anthropic='a', gemini='g').provedor, '')
    # a resolução acontece em carrega(); aqui conferimos a regra
    for anthropic, gemini_, esperado in (('a', '', 'claude'), ('', 'g', 'gemini'),
                                         ('a', 'g', 'claude'), ('', '', 'claude')):
        c = Config(anthropic=anthropic, gemini=gemini_)
        c.provedor = 'gemini' if (c.gemini and not c.anthropic) else 'claude'
        igual(c.provedor, esperado, f'{anthropic!r}/{gemini_!r}')


@teste('config: a chave e o modelo do cérebro seguem o provedor')
def _():
    c = cfg_teste(provedor='gemini', gemini='g', anthropic='a',
                  modelo_gemini='gemini-x', modelo='claude-y')
    igual(c.chave_do_cerebro, 'g')
    igual(c.modelo_do_cerebro, 'gemini-x')
    c.provedor = 'claude'
    igual(c.chave_do_cerebro, 'a')
    igual(c.modelo_do_cerebro, 'claude-y')


@teste('diagnóstico: chave do Gemini recusada não passa por válida')
def _():
    import jarvis as J
    import nucleo.modelos as M
    original = M.modelos_gemini, M.gemini
    M.gemini = lambda chave='': object()
    def recusa(cli):
        raise RuntimeError('400 INVALID_ARGUMENT: API key not valid')
    M.modelos_gemini = recusa
    try:
        bem, detalhe = J._testa_chave(cfg_teste(provedor='gemini', gemini='errada'))
        igual(bem, False, 'chave falsa passou por válida')
        verdade('RECUSOU' in detalhe)
        verdade('aistudio.google.com' in detalhe, 'precisa dizer onde pegar a certa')
    finally:
        M.modelos_gemini, M.gemini = original


# ══ ferramentas de negócio ══════════════════════════════════════════
@teste('funil: traduz "negociação" para os estados certos do banco')
def _():
    from ferramentas import funil as F
    igual(F._fase('negociação'), ['respondeu', 'quer_demo', 'demo_pronta'])
    igual(F._fase('negociacao'), F._fase('negociação'))
    igual(F._fase('fechados'), ['fechado'])
    igual(F._fase('inventada'), [])


@teste('funil: responde do banco de verdade, e avisa quando não acha')
def _():
    import sqlite3

    from ferramentas import funil as F
    banco = TMP / 'funil.db'
    cx = sqlite3.connect(banco)
    cx.executescript("""
      CREATE TABLE leads (id INTEGER PRIMARY KEY, nome TEXT, cidade TEXT, estado TEXT,
        telefone TEXT, instagram TEXT, avaliacoes INT, pontuacao INT);
      CREATE TABLE mensagens (id INTEGER PRIMARY KEY, lead_id INT, tipo TEXT,
        situacao TEXT, texto TEXT);
      CREATE TABLE demos (id INTEGER PRIMARY KEY, lead_id INT, url TEXT, situacao TEXT);
      INSERT INTO leads VALUES (1,'Grão Dourado','Natal','quer_demo','84','graodourado',120,9),
                              (2,'Mega Express','SRN','respondeu','89','mega',80,7),
                              (3,'Boteco X','Natal','novo','84','',10,4);
    """)
    cx.commit(); cx.close()
    ctx = Contexto(cfg=cfg_teste(funil_db=str(banco)))
    r = F.clientes_na_fase(F.FaseArgs(fase='negociacao'), ctx)
    verdade('Grão Dourado' in r and 'Mega Express' in r)
    verdade('Boteco X' not in r, 'lead novo não está em negociação')
    verdade('2 em negociacao' in r)
    resumo = F.clientes_resumo(F.NadaArgs(), ctx)
    verdade('3 clientes no funil' in resumo)
    vazio = F.clientes_resumo(F.NadaArgs(), Contexto(cfg=cfg_teste(funil_db='/nao/existe.db')))
    verdade('não achei' in vazio, 'tem que dizer o que fazer, não só falhar')


@teste('whatsapp: lê o carimbo de hora e autor da mensagem')
def _():
    from ferramentas import whatsapp as W
    quando, autor = W._quebra_meta('[20:31, 02/10/2026] Fulano: ')
    igual(autor, 'Fulano')
    verdade(quando > 0)
    igual(W._quebra_meta('qualquer coisa'), (0, ''))


@teste('whatsapp: a instrução proíbe contar orçamento como venda')
def _():
    from ferramentas import whatsapp as W
    t = W.EXTRACAO.lower()
    for exigido in ('orçamento', 'vou pensar', 'na dúvida, não conte'):
        verdade(exigido in t, f'a instrução precisa barrar "{exigido}"')
    igual(REGISTRO['whatsapp_enviar'].nivel, PERIGO)


# ══ slides ══════════════════════════════════════════════════════════
@teste('slides: monta um .pptx válido, com capa, conteúdo e fechamento')
def _():
    from pptx import Presentation

    from ferramentas.slides import Apresentacao, Slide, monta
    a = Apresentacao(titulo='Teste', subtitulo='sub',
                     slides=[Slide(titulo='Um', pontos=['a', 'b'], nota='falar disso'),
                             Slide(titulo='Dois', pontos=['c'])],
                     fechamento='fim')
    destino = monta(a, TMP / 'x.pptx', 'escuro')
    verdade(zipfile.is_zipfile(destino), 'pptx tem que ser um zip válido')
    p = Presentation(str(destino))
    igual(len(p.slides), 4, 'capa + 2 + fechamento')
    textos = ' '.join(sh.text_frame.text for s in p.slides for sh in s.shapes
                      if sh.has_text_frame)
    verdade('Um' in textos and 'Dois' in textos and 'fim' in textos)
    igual(p.slides[1].notes_slide.notes_text_frame.text, 'falar disso')


@teste('slides: o texto nunca sai da área do slide')
def _():
    from pptx.util import Inches

    from ferramentas.slides import ALTURA, LARGURA, Apresentacao, Slide, monta
    from pptx import Presentation
    a = Apresentacao(titulo='T', slides=[Slide(titulo='X', pontos=[f'ponto {i}'
                                                                  for i in range(6)])])
    p = Presentation(str(monta(a, TMP / 'y.pptx')))
    for i, s in enumerate(p.slides):
        for sh in s.shapes:
            verdade(sh.left >= 0 and sh.top >= 0, f'slide {i}: forma fora pela esquerda/topo')
            verdade(sh.left + sh.width <= LARGURA + Inches(0.1),
                    f'slide {i}: forma passa da largura')
            verdade(sh.top + sh.height <= ALTURA + Inches(0.1),
                    f'slide {i}: forma passa da altura')


@teste('slides: todo tema passa no contraste AA')
def _():
    from ferramentas.slides import PALETAS

    def lum(h):
        c = [int(h[i:i + 2], 16) / 255 for i in (0, 2, 4)]
        f = lambda v: v / 12.92 if v <= 0.03928 else ((v + 0.055) / 1.055) ** 2.4
        r, g, b = [f(x) for x in c]
        return 0.2126 * r + 0.7152 * g + 0.0722 * b

    for nome, p in PALETAS.items():
        for campo in ('texto', 'fraco', 'acento'):
            l1, l2 = lum(p[campo]), lum(p['fundo'])
            ct = (max(l1, l2) + 0.05) / (min(l1, l2) + 0.05)
            verdade(ct >= 4.5, f'tema {nome}, {campo}: {ct:.2f}:1')


@teste('config: chave salva pelo Bloco de Notas (com BOM) continua sendo lida')
def _():
    import os

    from nucleo.config import _carrega_env
    env = TMP / 'com-bom.env'
    # É assim que o Bloco de Notas do Windows salva em UTF-8.
    env.write_bytes('\ufeffANTHROPIC_API_KEY=sk-ant-teste-123\n'
                    'OPENAI_API_KEY="com-aspas"\n'.encode('utf-8'))
    for k in ('ANTHROPIC_API_KEY', 'OPENAI_API_KEY'):
        os.environ.pop(k, None)
    try:
        _carrega_env(env)
        igual(os.environ.get('ANTHROPIC_API_KEY'), 'sk-ant-teste-123',
              'o BOM comeu o nome da chave')
        igual(os.environ.get('OPENAI_API_KEY'), 'com-aspas', 'as aspas vazaram')
    finally:
        for k in ('ANTHROPIC_API_KEY', 'OPENAI_API_KEY'):
            os.environ.pop(k, None)


@teste('diagnóstico: token do claude.ai é reconhecido antes de ir à rede')
def _():
    import jarvis as J
    bem, detalhe = J._testa_chave(cfg_teste(anthropic='sk-ant-usr-1abcdef'))
    igual(bem, False)
    verdade('token do claude.ai' in detalhe, f'não identificou o token: {detalhe}')
    verdade('sk-ant-api03-' in detalhe, 'precisa dizer com o que a certa começa')
    bem, detalhe = J._testa_chave(cfg_teste(anthropic='minha-chave'))
    igual(bem, False)
    verdade('sk-ant-' in detalhe)


@teste('diagnóstico: chave recusada pela API é reportada como recusada')
def _():
    import types as _t

    import jarvis as J

    class Recusa:
        class messages:
            @staticmethod
            def count_tokens(**kw):
                raise RuntimeError('Error code: 401 - authentication_error: invalid x-api-key')

    import anthropic
    original = anthropic.Anthropic
    anthropic.Anthropic = lambda **kw: Recusa()
    try:
        bem, detalhe = J._testa_chave(cfg_teste(anthropic='sk-ant-errada'))
        igual(bem, False)
        verdade('RECUSOU' in detalhe, f'mensagem inútil: {detalhe}')
        verdade('console.anthropic.com' in detalhe, 'precisa dizer onde resolver')
    finally:
        anthropic.Anthropic = original


@teste('ouvido: sem motor local, a chave do Gemini salva o microfone')
def _():
    import importlib
    import sys as _sys

    from nucleo.ouvido import Ouvido
    pasta = TMP / 'mod2'
    pasta.mkdir(exist_ok=True)
    (pasta / 'faster_whisper.py').write_text(
        'raise ImportError("DLL load failed while importing resampler")')
    (pasta / 'sounddevice.py').write_text('')
    _sys.path.insert(0, str(pasta))
    for m in ('faster_whisper', 'sounddevice'):
        _sys.modules.pop(m, None)
    importlib.invalidate_caches()
    try:
        com = cfg_teste(gemini='AQ.chave')
        ok, como = Ouvido.checa_com(com)
        igual(ok, True, 'com a chave do Gemini o microfone deveria funcionar')
        verdade('Gemini' in como)
        verdade('sair da máquina' in como, 'a troca de privacidade precisa ficar dita')

        sem = cfg_teste(gemini='')
        ok2, _ = Ouvido.checa_com(sem)
        igual(ok2, False, 'sem motor e sem chave não há como ouvir')
    finally:
        _sys.path.remove(str(pasta))
        for m in ('faster_whisper', 'sounddevice'):
            _sys.modules.pop(m, None)
        importlib.invalidate_caches()


@teste('ouvido: o áudio vira um WAV de 16 bits que o Gemini aceita')
def _():
    import wave
    import io

    import numpy as np

    from nucleo.ouvido import Ouvido, TAXA
    som = (np.sin(np.linspace(0, 400, TAXA)) * 0.5).astype(np.float32)
    dados = Ouvido.para_wav(som)
    verdade(dados.startswith(b'RIFF'), 'não é WAV')
    with wave.open(io.BytesIO(dados)) as w:
        igual(w.getnchannels(), 1)
        igual(w.getsampwidth(), 2, 'tem que ser 16 bits')
        igual(w.getframerate(), TAXA)
        igual(w.getnframes(), TAXA)
    # o estouro é cortado, não dá a volta virando ruído
    alto = np.array([2.0, -2.0, 0.0], dtype=np.float32)
    quadros = np.frombuffer(Ouvido.para_wav(alto)[44:], dtype=np.int16)
    igual(list(quadros), [32767, -32767, 0])


@teste('ouvido: motor escolhido uma vez, e respeitando o config')
def _():
    from nucleo.ouvido import Ouvido
    o = Ouvido(cfg_teste(motor_escuta='gemini', gemini='x'))
    igual(o.motor(), 'gemini')
    o2 = Ouvido(cfg_teste(motor_escuta='local'))
    igual(o2.motor(), 'local', 'quem pediu local não pode ser mandado para a nuvem')


@teste('ouvido: pacote instalado que quebra não vira "falta instalar"')
def _():
    # O caso real: faster-whisper instalado e falhando ao carregar (DLL,
    # versão de Python). Dizer "falta instalar" manda a pessoa rodar um pip
    # que responde "already satisfied" — e o problema continua de pé.
    import importlib
    import sys as _sys
    pasta = TMP / 'modulos'
    pasta.mkdir(exist_ok=True)
    (pasta / 'faster_whisper.py').write_text(
        'raise ImportError("DLL load failed while importing _ext")')
    (pasta / 'sounddevice.py').write_text('')
    _sys.path.insert(0, str(pasta))
    for m in ('faster_whisper', 'sounddevice'):
        _sys.modules.pop(m, None)
    importlib.invalidate_caches()
    try:
        from nucleo.ouvido import Ouvido
        ok, aviso = Ouvido.checa()
        igual(ok, False)
        verdade('não carrega' in aviso, f'mensagem errada: {aviso}')
        verdade('DLL load failed' in aviso, 'o erro real precisa aparecer')
        verdade('falta instalar' not in aviso, 'mandou instalar o que já está lá')
    finally:
        _sys.path.remove(str(pasta))
        for m in ('faster_whisper', 'sounddevice'):
            _sys.modules.pop(m, None)
        importlib.invalidate_caches()


@teste('ouvido: pacote que realmente falta continua pedindo instalação')
def _():
    from nucleo.ouvido import Ouvido
    ok, aviso = Ouvido.checa()
    if ok:
        return                       # a máquina tem tudo; nada a verificar
    verdade('falta instalar' in aviso or 'não carrega' in aviso)


# ══ navegador e voz ═════════════════════════════════════════════════
@teste('navegador: entende site, URL e frase de busca')
def _():
    from ferramentas.navegador import _url
    igual(_url('instagram.com'), 'https://instagram.com')
    igual(_url('https://x.com/a'), 'https://x.com/a')
    verdade(_url('quanto custa um site').startswith('https://www.google.com/search?q='))
    igual(_url('  globo.com.br  '), 'https://globo.com.br')


@teste('voz: tira da fala o que é para o olho')
def _():
    igual(limpa_para_fala('Veja **isto** em `config.toml`'), 'Veja isto em config.toml')
    verdade('o link está na tela' in limpa_para_fala('abre https://exemplo.com/a?b=1'))
    verdade('o código está na tela' in limpa_para_fala('assim:\n```py\nx=1\n```'))
    igual(limpa_para_fala('- um\n- dois'), 'um\ndois')


@teste('voz: cale() esvazia a fila em vez de só parar o que toca')
def _():
    from nucleo.voz import Voz
    import os
    os.environ['JARVIS_SEM_VOZ'] = '1'
    v = Voz(cfg_teste())
    for i in range(20):
        v.fala(f'frase {i}')
    v.cale()
    time.sleep(0.2)
    verdade(v._fila.qsize() == 0, 'sobrou fala na fila depois do cale')


# ══ corrida ═════════════════════════════════════════════════════════
def main() -> int:
    ok = falhas = 0
    print()
    for nome, f in CASOS:
        barulho = io.StringIO()
        try:
            with contextlib.redirect_stdout(barulho):
                f()
            ok += 1
            print(f'  ✓ {nome}')
        except Exception as e:
            falhas += 1
            print(f'  ✗ {nome}\n      {type(e).__name__}: {e}')
            if barulho.getvalue():
                for linha in barulho.getvalue().splitlines()[-6:]:
                    print(f'      | {linha}')
            if '-v' in sys.argv:
                traceback.print_exc()
    shutil.rmtree(TMP, ignore_errors=True)
    print(f'\n  {ok + falhas} testes · {ok} passando · {falhas} falhando\n')
    return 1 if falhas else 0


if __name__ == '__main__':
    raise SystemExit(main())
